"""Los Excel que exporta el MIR: archivo para corregir, informe e instructivo.

Rediseñados el 26-09-2026 con el responsable funcional. Estos tests fijan lo
que se acordó, sin base de datos: cada Excel se arma a partir de datos ya
leídos.
"""

import io
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook

from runac.services import informe_errores_service as informes
from runac.services import instructivo_service as instructivos

_MOTOR = Path(__file__).resolve().parents[1] / "services" / "motor"
if str(_MOTOR) not in sys.path:
    sys.path.insert(0, str(_MOTOR))

import en_palabras  # noqa: E402  # pylint: disable=wrong-import-position

GENERADO = datetime(2026, 9, 26, 11, 40)


# ---------------------------------------------------------------------------
# En palabras
# ---------------------------------------------------------------------------


def test_las_reglas_se_dicen_en_palabras_y_nombran_los_campos_por_su_titulo():
    titulos = {"presenta_alguna_discapacidad": "¿Presenta alguna discapacidad?"}
    regla = {
        "tipo": "OBLIGATORIO_SI",
        "parametros": '{"operador": "IGUAL", "campo_condicion": "presenta_alguna_discapacidad", "valor_condicion": "Si"}',
    }
    assert (
        en_palabras.regla_en_palabras(regla, titulos)
        == "Obligatorio si «¿Presenta alguna discapacidad?» es «Si»."
    )
    hoy = {
        "tipo": "COMPARAR_VALOR",
        "parametros": {"valor": "HOY", "operador": "MENOR_IGUAL"},
    }
    assert en_palabras.regla_en_palabras(hoy, {}) == "No puede ser posterior a hoy."


def test_el_problema_no_suena_a_maquina_ni_repite_el_valor():
    # Así lo guardaba el motor, y así se mostraba hasta el 26-09-2026.
    assert (
        en_palabras.problema_en_palabras(
            "COMPARAR_VALOR",
            "El valor no cumple la condición: debe ser menor igual hoy.",
        )
        == "Fecha posterior a hoy: tiene que ser de hoy o anterior."
    )
    texto = en_palabras.problema_en_palabras(
        "OBLIGATORIO_SI",
        "El campo es obligatorio cuando «¿Presenta alguna discapacidad?» igual «Si».",
    )
    assert (
        texto
        == "Falta el dato: es obligatorio porque «¿Presenta alguna discapacidad?» es «Sí»."
    )
    lista = en_palabras.problema_en_palabras(
        "FUERA_DE_CATALOGO", "El valor no está entre los admitidos.", ["DNI", "Otro"]
    )
    assert "DNI, Otro" in lista


def test_lo_que_no_se_reconoce_queda_como_lo_dijo_el_motor():
    assert (
        en_palabras.problema_en_palabras("ALGO_NUEVO", "Texto del motor.")
        == "Texto del motor."
    )


def test_obligatorio_segun_otro_dato():
    assert (
        en_palabras.obligatoriedad({"obligatorio": 0}, [{"tipo": "OBLIGATORIO_SI"}])
        == en_palabras.SEGUN_OTRO
    )


# ---------------------------------------------------------------------------
# Informe de la importación
# ---------------------------------------------------------------------------


def _cabecera(**extra):
    return {
        "id": 11,
        "nombre_archivo": "MPI_2026_T1_Chubut.xlsx",
        "estado": "VALIDA",
        "usuario": "operador",
        "filas_leidas": 30,
        "filas_incorporadas": 30,
        "iniciada_el": datetime(2026, 9, 25, 19, 50),
        "archivo_codigo": "MPI",
        "archivo_nombre": "Nómina MPI",
        "version": 1,
        "que_es_una_fila": "Niño, niña o adolescente",
        "jurisdiccion": "Chubut",
        "periodo": "2026_T1",
        **extra,
    }


def _regla(**extra):
    return {
        "nombre_hoja": "MPI",
        "numero_fila": 12,
        "columna": "T",
        "nombre_campo": "Edad",
        "campo_id": 201,
        "severidad": "ADVERTENCIA",
        "codigo": "RANGO",
        "valor_encontrado": "107",
        "descripcion": "El valor 107 es mayor que el máximo esperado (17).",
        "identificador_registro": "9 · Paz · Ana",
        "resuelta": 0,
        "catalogo_id": None,
        **extra,
    }


def _informe(datos):
    return load_workbook(io.BytesIO(informes.armar_informe(datos, GENERADO)))


def test_el_informe_abre_en_los_problemas_y_termina_en_el_resumen():
    datos = {
        "cabecera": _cabecera(),
        "reglas": [_regla()],
        "archivo": [],
        "cambios": {},
        "opciones": {},
    }
    wb = _informe(datos)
    assert wb.sheetnames == ["Problemas", "Resumen"]
    assert wb.active.title == "Problemas"
    ws = wb["Problemas"]
    # El título de la identificación dice qué es una fila de este archivo.
    assert ws["D1"].value == "Niño, niña o adolescente"
    assert ws["G2"].value == "Advertencia"
    assert ws["H2"].value.startswith("Fuera de lo esperable (hasta 17)")


def test_cada_problema_en_su_fila_sin_pisar_los_titulos():
    reglas = [_regla(numero_fila=f) for f in (5, 6, 7)]
    ws = _informe(
        {
            "cabecera": _cabecera(),
            "reglas": reglas,
            "archivo": [],
            "cambios": {},
            "opciones": {},
        }
    )["Problemas"]
    assert ws["A1"].value == "Hoja"
    assert [ws.cell(row=r, column=2).value for r in (2, 3, 4)] == [5, 6, 7]


def test_lo_corregido_en_el_sistema_dice_por_quien_y_cuando():
    cambio = {
        "valor_anterior": "107",
        "valor_nuevo": "12",
        "usuario": "operador",
        "fecha": datetime(2026, 9, 25, 20, 45),
    }
    datos = {
        "cabecera": _cabecera(),
        "reglas": [_regla(resuelta=1)],
        "archivo": [],
        "cambios": {(12, 201): cambio},
        "opciones": {},
    }
    ws = _informe(datos)["Problemas"]
    assert ws["I2"].value == "Corregido"
    assert ws["J2"].value == "107 → 12 · operador · 25/09/2026 20:45"


def test_el_estado_es_un_desplegable_con_las_cuatro_opciones():
    ws = _informe(
        {
            "cabecera": _cabecera(),
            "reglas": [_regla()],
            "archivo": [],
            "cambios": {},
            "opciones": {},
        }
    )["Problemas"]
    formulas = [dv.formula1 for dv in ws.data_validations.dataValidation]
    assert '"Pendiente,Corregido,Está bien así,A consultar"' in formulas


def test_el_informe_tampoco_devuelve_formulas():
    datos = {
        "cabecera": _cabecera(),
        "reglas": [_regla(valor_encontrado='=HYPERLINK("x")')],
        "archivo": [],
        "cambios": {},
        "opciones": {},
    }
    ws = _informe(datos)["Problemas"]
    assert ws["F2"].data_type == "s"


def test_ninguna_solapa_tiene_color():
    """Una solapa pintada llama más la atención que la activa (26-09-2026)."""
    wb = _informe(
        {
            "cabecera": _cabecera(),
            "reglas": [_regla()],
            "archivo": [],
            "cambios": {},
            "opciones": {},
        }
    )
    assert all(ws.sheet_properties.tabColor is None for ws in wb.worksheets)


# ---------------------------------------------------------------------------
# Archivo para corregir
# ---------------------------------------------------------------------------


def _definicion():
    campo = {
        "nombre": "edad",
        "titulo_esperado": "Edad",
        "orden": 1,
        "tipo_dato": "ENTERO",
        "longitud_maxima": None,
        "obligatorio": 1,
        "ayuda": None,
        "catalogo": None,
        "reglas": [],
        "opciones": [],
    }
    return {
        "codigo": "MPI",
        "titulos": {"edad": "Edad"},
        "hojas": [
            {
                "nombre_esperado": "MPI",
                "fila_encabezados": 2,
                "referencia": 0,
                "campos": [campo],
            },
            {
                "nombre_esperado": "Prov_Dto_Localidad",
                "fila_encabezados": 1,
                "referencia": 1,
                "campos": [],
            },
        ],
    }


def _subido():
    wb = Workbook()
    wb.active.title = "MPI"
    wb["MPI"]["A2"] = "Edad *"
    wb["MPI"]["A12"] = 107
    wb.create_sheet("Prov_Dto_Localidad")
    wb.create_sheet("LISTAS")
    wb.create_sheet("ADVERTENCIAS_ESPERADAS")
    return wb


def _para_corregir(reglas):
    datos = {
        "cabecera": _cabecera(),
        "reglas": reglas,
        "archivo": [],
        "cambios": {},
        "opciones": {},
    }
    return load_workbook(
        io.BytesIO(
            informes.armar_archivo_para_corregir(
                _subido(), datos, _definicion(), GENERADO
            )
        )
    )


def test_quedan_las_hojas_de_datos_listas_y_resumen():
    wb = _para_corregir([_regla(columna="A")])
    assert wb.sheetnames == ["MPI", "Prov_Dto_Localidad", "Listas", "Resumen"]
    # La de referencia, oculta: parecía que había que completarla (#87).
    assert wb["Prov_Dto_Localidad"].sheet_state == "hidden"


def test_la_celda_con_problema_se_marca_y_explica_que_hacer():
    ws = _para_corregir([_regla(columna="A")])["MPI"]
    assert "Advertencia" in ws["A12"].comment.text
    assert "Fuera de lo esperable" in ws["A12"].comment.text
    # El título explica qué va: lo que antes decía la hoja de instrucciones.
    assert "Obligatorio" in ws["A2"].comment.text


def test_abre_en_el_primer_problema():
    wb = _para_corregir([_regla(columna="A")])
    assert wb.active.title == "MPI"
    assert wb["MPI"].sheet_view.selection[0].activeCell == "A12"


def test_sin_problemas_abre_en_el_resumen():
    assert _para_corregir([]).active.title == "Resumen"


# ---------------------------------------------------------------------------
# Instructivo
# ---------------------------------------------------------------------------


def test_el_instructivo_tiene_general_y_campos_y_las_listas_no_se_ven():
    definicion = _definicion()
    definicion["hojas"][0]["campos"][0].update(catalogo="si_no", opciones=["Sí", "No"])
    definicion.update(descripcion="Nómina MPI", version=1)
    wb = load_workbook(io.BytesIO(instructivos.armar(definicion, "2026_T1", GENERADO)))
    assert [(ws.title, ws.sheet_state) for ws in wb.worksheets] == [
        ("General", "visible"),
        ("Campos", "visible"),
        ("listas", "veryHidden"),
    ]
    campos = wb["Campos"]
    # Una fila por campo, y la hoja de referencia no aparece.
    assert campos.max_row == 2
    assert campos["G2"].value == "2 opciones ▾"
    assert any(
        dv.formula1.startswith("listas!")
        for dv in campos.data_validations.dataValidation
    )
