"""El instructivo de un archivo: una fila por campo, para bajar junto con la plantilla.

Se descarga aparte de la plantilla (pedido del responsable funcional,
26-09-2026). Antes la explicación iba adentro de la plantilla, en una hoja
INSTRUCCIONES de texto corrido que no decía el tipo, el largo ni la lista de
cada campo.

Lee la definición con el mismo `leer_definicion` que la plantilla y la explica
con las mismas palabras (`motor/en_palabras.py`): lo que dice el instructivo y
lo que dice el globo de cada título de la plantilla no pueden diferir.

Dos solapas: «General», que explica cómo se lee, y «Campos». Las opciones de
cada lista se ven en un desplegable, alimentado por una hoja que no se ve. Excel
no deja mostrar un desplegable sin dejar elegir: si alguien elige, la celda
cambia de texto, y en un instructivo eso no hace daño.
"""

from __future__ import annotations

import io
from datetime import datetime

from django.db import connection
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from runac.services import excel_estilo as e
from runac.services import plantillas_service

plantilla = plantillas_service.motor_plantilla
en_palabras = plantilla.en_palabras


def nombre_de_archivo(codigo: str, periodo: str) -> str:
    return f"{codigo}_{periodo}_INSTRUCTIVO.xlsx"


def generar(codigo: str, periodo: str, usuario: str = "") -> bytes:
    with connection.cursor() as cursor:
        archivo = plantilla.leer_definicion(
            plantillas_service._CursorConNombres(cursor), codigo
        )
    return armar(archivo, periodo, datetime.now(), usuario)


def armar(archivo: dict, periodo: str, generado: datetime, usuario: str = "") -> bytes:
    """El libro, a partir de la definición ya leída. Separado para probarlo sin base."""
    hojas = [h for h in archivo["hojas"] if not h.get("referencia")]
    campos = [(h, c) for h in hojas for c in h["campos"]]
    titulos = archivo.get("titulos", {})
    codigo = archivo["codigo"]
    pie = f"MIR · Instructivo {codigo} · {periodo}"

    wb = Workbook()
    _general(wb.active, archivo, periodo, generado, campos, titulos, pie, usuario)

    ws = wb.create_sheet("Campos")
    e.encabezado(
        ws,
        [
            "Hoja",
            "Grupo",
            "Columna",
            "Campo",
            "Obligatorio",
            "Qué va",
            "Valores posibles",
            "Controles",
            "Para qué sirve",
        ],
        [12, 22, 8, 32, 17, 22, 22, 46, 50],
    )
    listas = wb.create_sheet("listas")
    listas.sheet_state = "veryHidden"
    columna_de = {}
    for hoja, c in campos:
        reglas = c.get("reglas", [])
        ob = en_palabras.obligatoriedad(c, reglas)
        cond = en_palabras.condiciones(reglas, titulos)
        ctrl = en_palabras.controles(reglas, titulos)
        opciones = c.get("opciones", [])
        ws.append(
            [
                hoja["nombre_esperado"],
                c.get("dimension") or "—",
                get_column_letter(c["orden"]),
                c["titulo_esperado"],
                "\n".join([ob] + cond),
                en_palabras.que_va(c),
                f"{len(opciones)} opciones ▾" if opciones else "",
                "\n".join(ctrl),
                (c.get("ayuda") or "").strip(),
            ]
        )
        n = ws.max_row
        for celda in ws[n]:
            e.cuerpo(celda)
        ws.cell(row=n, column=3).alignment = Alignment(
            horizontal="center", vertical="top"
        )
        # El campo con el color de su título en la plantilla: así se lo reconoce.
        campo = ws.cell(row=n, column=4)
        campo.fill = e.relleno(
            e.VERDE_OBLIGATORIO if c.get("obligatorio") else e.VERDE_OPCIONAL
        )
        campo.font = e.fuente(size=10, bold=True, color=e.TINTA)
        if ob == en_palabras.SEGUN_OTRO:
            ws.cell(row=n, column=5).font = e.fuente(
                size=10, bold=True, color=e.ADVERTENCIA["tin"]
            )
        if opciones:
            clave = c.get("catalogo") or c["titulo_esperado"]
            if clave not in columna_de:
                col = len(columna_de) + 1
                columna_de[clave] = get_column_letter(col)
                for i, o in enumerate(opciones, start=1):
                    listas.cell(row=i, column=col, value=o)
            letra = columna_de[clave]
            dv = DataValidation(
                type="list",
                formula1=f"listas!${letra}$1:${letra}${len(opciones)}",
                allow_blank=True,
                showErrorMessage=False,
                showInputMessage=True,
                promptTitle="Valores posibles",
                prompt="Abrir el desplegable para ver las opciones.",
            )
            ws.add_data_validation(dv)
            dv.add(ws.cell(row=n, column=7))
            ws.cell(row=n, column=7).font = e.fuente(size=10, bold=True, color=e.TEAL)
        e.alto_para(
            ws,
            n,
            [
                ws.cell(row=n, column=5).value,
                ws.cell(row=n, column=8).value,
                c.get("ayuda"),
            ],
        )
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = f"A1:I{ws.max_row}"
    e.para_imprimir(ws, pie, repetir="1:1")

    wb.active = 0
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _general(ws, archivo, periodo, generado, campos, titulos, pie, usuario=""):
    ws.title = "General"
    codigo = archivo["codigo"]
    e.portada(
        ws,
        f'Instructivo · {archivo.get("descripcion") or codigo}',
        f'{codigo} · período {periodo} · estructura v{archivo.get("version")} · '
        f"generado el {e.ahora_legible(generado)}"
        + (f" por {usuario}" if usuario else ""),
        "",
        ancho=2,
    )
    ws.column_dimensions["B"].width = 26
    ws.column_dimensions["C"].width = 86
    bloques = [
        (
            "Para qué es",
            "Explica cada columna de la plantilla: qué va, si es obligatoria, qué valores admite y qué "
            "controla el sistema al importar. La plantilla y este instructivo se descargan por separado.",
        ),
        (
            "Cómo se lee",
            "La solapa «Campos» tiene una fila por columna de la plantilla, en el mismo orden. En "
            "«Valores posibles», el desplegable muestra las opciones de la lista: es para consultarlas, "
            "no hace falta elegir nada.",
        ),
        (
            "Nombre del archivo",
            f"El archivo se sube con el nombre {codigo}_{periodo}_<jurisdicción>.xlsx; por "
            f"ejemplo, {codigo}_{periodo}_Chubut.xlsx.",
        ),
        ("Colores de los títulos", None),
        (
            "Bloqueante y advertencia",
            "Si un control es bloqueante y no se cumple, no se incorpora ninguna fila del "
            "archivo hasta corregirlo. Si es una advertencia (en «Controles» dice «Aviso»), el "
            "archivo entra igual y el dato se puede revisar o corregir dentro del sistema.",
        ),
        (
            "Qué no hay que tocar",
            "Los títulos de las columnas, su orden y los nombres de las hojas: el sistema los "
            "busca tal cual. Las filas de datos van seguidas, sin filas vacías en el medio.",
        ),
    ]
    fila = 7
    for titulo, texto in bloques:
        ws.cell(row=fila, column=2, value=titulo).font = e.fuente(
            size=11, bold=True, color=e.TINTA
        )
        ws.cell(row=fila, column=2).alignment = Alignment(vertical="top")
        if texto:
            c = ws.cell(row=fila, column=3, value=texto)
            c.font = e.fuente(size=10, color=e.TINTA)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            e.alto_para(ws, fila, [texto], 95)
            fila += 2
            continue
        # La jerarquía de verdes de la plantilla (28-09-2026).
        for color, borde, que, tinta in (
            (
                e.VERDE_GRUPO,
                e.AMBAR,
                "Verde oscuro: grupo de columnas (el título que abarca a varias).",
                "FFFFFF",
            ),
            (
                e.VERDE_OBLIGATORIO,
                e.TEAL,
                "Verde: dato obligatorio (el título lleva *).",
                e.TINTA,
            ),
            (e.VERDE_OPCIONAL, "9DCECB", "Verde clarito: dato opcional.", e.TINTA),
        ):
            c = ws.cell(row=fila, column=3, value=que)
            c.fill = e.relleno(color)
            c.font = e.fuente(size=10, bold=True, color=tinta)
            c.border = Border(left=Side(style="thick", color=borde))
            fila += 1
        nota = (
            "Algunos datos opcionales pasan a ser obligatorios según otra respuesta: lo dice la columna "
            "«Obligatorio» de «Campos». Al pasar el mouse por un título de la plantilla se ve lo mismo que en "
            "este instructivo."
        )
        c = ws.cell(row=fila, column=3, value=nota)
        c.font = e.fuente(size=10, color=e.TINTA2)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        e.alto_para(ws, fila, [nota], 95)
        fila += 2

    cuenta = {
        en_palabras.OBLIGATORIO: 0,
        en_palabras.SEGUN_OTRO: 0,
        en_palabras.OPCIONAL: 0,
    }
    for _, c in campos:
        cuenta[en_palabras.obligatoriedad(c, c.get("reglas", []))] += 1
    ws.cell(row=fila, column=2, value="En números").font = e.fuente(
        size=11, bold=True, color=e.TINTA
    )
    ws.cell(
        row=fila,
        column=3,
        value=f"{len(campos)} campos: {cuenta[en_palabras.OBLIGATORIO]} obligatorios, "
        f"{cuenta[en_palabras.SEGUN_OTRO]} obligatorios según otro dato y {cuenta[en_palabras.OPCIONAL]} opcionales.",
    ).font = e.fuente(size=10, color=e.TINTA)
    e.para_imprimir(ws, pie)
