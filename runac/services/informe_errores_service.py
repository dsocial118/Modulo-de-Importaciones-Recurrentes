"""Informes de errores que el operador se lleva para corregir.

El operador corrige **en el Excel**, no en la pantalla. Por eso hay dos salidas,
y la segunda es la que realmente le sirve:

1. `planilla_de_errores`: un Excel nuevo con una hoja por cada hoja del archivo
   y una fila por error, con su celda, su valor y qué hay que corregir.

2. `archivo_marcado`: **el mismo archivo que subió**, con las celdas
   problemáticas pintadas y un comentario en cada una. Lo abre, ve exactamente
   qué corregir y dónde, corrige y lo vuelve a importar.

Ninguna de las dos toca la base de datos más allá de leerla.
"""

from __future__ import annotations

import io
import os
from typing import Any

from django.db import connection
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# Los mismos colores que usa el prototipo en pantalla, para que el Excel y la
# web se lean como una sola cosa.
AZUL = "1F4E79"
ROJO_SUAVE = "F8D7DA"
AMARILLO_SUAVE = "FFF3CD"
GRIS = "F2F2F2"

COLOR_POR_SEVERIDAD = {"BLOQUEANTE": ROJO_SUAVE, "ADVERTENCIA": AMARILLO_SUAVE}


def _filas(cursor) -> list[dict[str, Any]]:
    columnas = [c[0] for c in cursor.description]
    return [dict(zip(columnas, fila)) for fila in cursor.fetchall()]


def datos_de_la_importacion(importacion_id: int) -> dict[str, Any]:
    """Cabecera, reglas incumplidas y problemas del archivo."""
    with connection.cursor() as cur:
        cur.execute(
            """
            SELECT i.id, i.nombre_archivo, i.ruta_archivo, i.estado,
                   i.filas_leidas, i.filas_incorporadas, i.bloqueantes, i.advertencias,
                   i.iniciada_el, a.codigo AS archivo_codigo, av.numero AS version,
                   j.nombre AS jurisdiccion, p.codigo AS periodo
            FROM runac_c2_importacion i
            LEFT JOIN runac_c1_archivo a ON a.id = i.archivo_id
            LEFT JOIN runac_c1_archivo_version av ON av.id = i.archivo_version_id
            LEFT JOIN runac_c2_presentacion s ON s.id = i.presentacion_id
            LEFT JOIN runac_c2_jurisdiccion j ON j.id = s.jurisdiccion_id
            LEFT JOIN runac_c2_periodo p ON p.id = s.periodo_id
            WHERE i.id = %s
            """,
            [importacion_id],
        )
        cabecera = _filas(cur)
        if not cabecera:
            return {}
        cabecera = cabecera[0]

        cur.execute(
            """
            SELECT r.nombre_hoja, r.numero_fila, r.columna, r.nombre_campo,
                   r.severidad, r.codigo, r.valor_encontrado, r.descripcion,
                   r.identificador_registro
            FROM runac_c2_reglas_incumplidas r
            JOIN runac_c2_importacion i ON i.id = r.importacion_id
            -- El orden de las hojas lo manda la Capa 1, no el alfabeto. Ordenado
            -- por nombre, el informe abria por CAD y seguia por CRC: el operador
            -- lo recorre contra su Excel, donde las hojas estan en otro orden.
            LEFT JOIN runac_c1_hoja h
                   ON h.archivo_version_id = i.archivo_version_id
                  AND h.nombre_esperado = r.nombre_hoja
            WHERE r.importacion_id = %s
            ORDER BY COALESCE(h.orden_procesamiento, 99), r.nombre_hoja,
                     r.numero_fila, r.columna
            """,
            [importacion_id],
        )
        reglas = _filas(cur)

        cur.execute(
            """
            SELECT tipo, hoja, numero_fila, esperado, encontrado, descripcion
            FROM runac_c2_errores_de_importacion
            WHERE importacion_id = %s ORDER BY id
            """,
            [importacion_id],
        )
        archivo = _filas(cur)

    return {"cabecera": cabecera, "reglas": reglas, "archivo": archivo}


# ---------------------------------------------------------------------------
# 1. Planilla de errores
# ---------------------------------------------------------------------------


def _inofensivo(valor):
    """Un texto que Excel no va a interpretar como fórmula.

    Todo lo que sale a estos informes viene de un archivo que subió una
    provincia. Excel trata como fórmula cualquier celda que empiece con `=`,
    `+`, `-` o `@`, así que un valor como `=HYPERLINK(...)` escrito en la
    planilla se ejecutaría en la máquina de quien abre el informe. Se marcan
    esas celdas como texto explícito y el valor se muestra tal cual vino: es un
    informe de lo que el archivo decía, y eso incluye lo que decía mal.
    """
    if isinstance(valor, str) and valor[:1] in ("=", "+", "-", "@"):
        return ("texto", valor)
    return (None, valor)


def _celda(hoja, fila: int, columna: int, valor):
    """Escribe una celda sin dejar que su contenido se convierta en fórmula."""
    formato, contenido = _inofensivo(valor)
    celda = hoja.cell(row=fila, column=columna, value=contenido)
    if formato == "texto":
        celda.data_type = "s"
        celda.quotePrefix = True
    return celda


def _escribir_encabezado(hoja, titulos: list[str]) -> None:
    for i, titulo in enumerate(titulos, start=1):
        celda = hoja.cell(row=1, column=i, value=titulo)
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor=AZUL)
    hoja.freeze_panes = "A2"


def planilla_de_errores(importacion_id: int) -> bytes:
    """Un Excel con una hoja por cada hoja del archivo y una fila por error."""
    datos = datos_de_la_importacion(importacion_id)
    if not datos:
        return b""

    cabecera = datos["cabecera"]
    libro = Workbook()
    libro.remove(libro.active)

    # Portada: de qué importación se trata y cómo se lee el informe.
    portada = libro.create_sheet("RESUMEN")
    portada.column_dimensions["A"].width = 28
    portada.column_dimensions["B"].width = 70
    resumen = [
        ("Archivo", cabecera["archivo_codigo"]),
        ("Nombre recibido", cabecera["nombre_archivo"]),
        ("Jurisdicción", cabecera["jurisdiccion"]),
        ("Período", cabecera["periodo"]),
        ("Estructura", f'versión {cabecera["version"]}'),
        ("Importado el", str(cabecera["iniciada_el"] or "")),
        ("Resultado", cabecera["estado"]),
        ("Filas leídas", cabecera["filas_leidas"]),
        ("Filas incorporadas", cabecera["filas_incorporadas"]),
        ("Errores bloqueantes", cabecera["bloqueantes"]),
        ("Advertencias", cabecera["advertencias"]),
    ]
    for i, (etiqueta, valor) in enumerate(resumen, start=1):
        portada.cell(row=i, column=1, value=etiqueta).font = Font(bold=True)
        _celda(portada, i, 2, valor)

    fila = len(resumen) + 2
    portada.cell(row=fila, column=1, value="Cómo leer este informe").font = Font(
        bold=True, size=12
    )
    ayuda = (
        "Hay una hoja por cada hoja del archivo que se importó. Cada fila del informe "
        "indica en qué fila y en qué columna del Excel original está el problema, qué "
        "valor se encontró y qué hay que corregir.\n\n"
        "Un error BLOQUEANTE impide la importación: mientras exista, no se incorpora "
        "ninguna fila del archivo. Una ADVERTENCIA no la impide y puede resolverse "
        "dentro del sistema."
    )
    celda = portada.cell(row=fila + 1, column=2, value=ayuda)
    celda.alignment = Alignment(wrap_text=True, vertical="top")
    portada.row_dimensions[fila + 1].height = 90

    # Una hoja por cada hoja del archivo original.
    por_hoja: dict[str, list[dict]] = {}
    for regla in datos["reglas"]:
        por_hoja.setdefault(regla["nombre_hoja"] or "SIN HOJA", []).append(regla)

    titulos = [
        "Fila",
        "Columna",
        "Campo",
        "Severidad",
        "Valor encontrado",
        "Qué hay que corregir",
        "Identificador",
        "Código",
    ]
    for nombre_hoja, errores in por_hoja.items():
        hoja = libro.create_sheet(nombre_hoja[:31])
        _escribir_encabezado(hoja, titulos)
        for i, e in enumerate(errores, start=2):
            valores = [
                e["numero_fila"],
                e["columna"],
                e["nombre_campo"],
                e["severidad"],
                (e["valor_encontrado"] or "")[:200],
                e["descripcion"],
                e["identificador_registro"],
                e["codigo"],
            ]
            for j, valor in enumerate(valores, start=1):
                celda = _celda(hoja, i, j, valor)
                celda.fill = PatternFill(
                    "solid", fgColor=COLOR_POR_SEVERIDAD.get(e["severidad"], GRIS)
                )
        for letra, ancho in zip("ABCDEFGH", (8, 10, 34, 14, 32, 62, 16, 22)):
            hoja.column_dimensions[letra].width = ancho

    # Los problemas del archivo entero van en su propia hoja: no tienen campo.
    if datos["archivo"]:
        hoja = libro.create_sheet("PROBLEMAS DEL ARCHIVO")
        _escribir_encabezado(
            hoja, ["Tipo", "Hoja", "Fila", "Se esperaba", "Se encontró", "Detalle"]
        )
        for i, e in enumerate(datos["archivo"], start=2):
            for j, valor in enumerate(
                [
                    e["tipo"],
                    e["hoja"],
                    e["numero_fila"],
                    e["esperado"],
                    e["encontrado"],
                    e["descripcion"],
                ],
                start=1,
            ):
                _celda(hoja, i, j, valor).fill = PatternFill(
                    "solid", fgColor=ROJO_SUAVE
                )
        for letra, ancho in zip("ABCDEF", (26, 22, 8, 26, 26, 62)):
            hoja.column_dimensions[letra].width = ancho

    if len(libro.sheetnames) == 1:  # sólo la portada: no hubo errores
        hoja = libro.create_sheet("SIN ERRORES")
        hoja["A1"] = "El archivo no registró errores ni advertencias."

    buffer = io.BytesIO()
    libro.save(buffer)
    return buffer.getvalue()


# ---------------------------------------------------------------------------
# 2. El archivo original, con las celdas marcadas
# ---------------------------------------------------------------------------


def archivo_marcado(importacion_id: int) -> tuple[bytes, str]:
    """Devuelve el Excel que subió el operador, con los errores señalados.

    Cada celda con problema queda pintada y con un comentario que explica qué
    corregir. Es la forma más directa de que el operador arregle el archivo:
    lo abre y ve exactamente dónde está cada cosa.

    Devuelve (contenido, nombre). Si no se conserva el archivo original,
    devuelve (b"", "").
    """
    datos = datos_de_la_importacion(importacion_id)
    if not datos:
        return b"", ""

    cabecera = datos["cabecera"]
    ruta = cabecera.get("ruta_archivo")
    if not ruta or not os.path.exists(ruta):
        return b"", ""

    libro = load_workbook(ruta)

    marcadas = 0
    for regla in datos["reglas"]:
        nombre_hoja = regla["nombre_hoja"]
        if not nombre_hoja or nombre_hoja not in libro.sheetnames:
            continue
        hoja = libro[nombre_hoja]
        columna = regla["columna"]
        fila = regla["numero_fila"]
        if not columna or not fila:
            continue
        try:
            celda = hoja[f"{columna}{fila}"]
        except (ValueError, KeyError):
            continue
        celda.fill = PatternFill(
            "solid", fgColor=COLOR_POR_SEVERIDAD.get(regla["severidad"], GRIS)
        )
        texto = (
            f'{regla["severidad"]}\n\n'
            f'{regla["descripcion"]}\n\n'
            f'Campo: {regla["nombre_campo"] or ""}'
        )
        comentario = Comment(texto, "RUNAC")
        comentario.width = 320
        comentario.height = 150
        celda.comment = comentario
        marcadas += 1

    # Una hoja al principio que explica qué es este archivo.
    guia = libro.create_sheet("CÓMO CORREGIR", 0)
    guia.column_dimensions["A"].width = 26
    guia.column_dimensions["B"].width = 78
    lineas = [
        ("Archivo", cabecera["archivo_codigo"]),
        ("Jurisdicción", cabecera["jurisdiccion"]),
        ("Período", cabecera["periodo"]),
        ("Celdas marcadas", marcadas),
        ("Errores bloqueantes", cabecera["bloqueantes"]),
        ("Advertencias", cabecera["advertencias"]),
    ]
    for i, (etiqueta, valor) in enumerate(lineas, start=1):
        guia.cell(row=i, column=1, value=etiqueta).font = Font(bold=True)
        _celda(guia, i, 2, valor)

    fila = len(lineas) + 2
    ayuda = (
        "Este es el mismo archivo que se importó, con las celdas problemáticas "
        "señaladas.\n\n"
        "· Las celdas en rojo tienen un error BLOQUEANTE: mientras existan, no se "
        "incorpora ninguna fila del archivo.\n"
        "· Las celdas en amarillo son ADVERTENCIAS: no impiden la importación.\n\n"
        "Al pasar el mouse sobre una celda marcada aparece la explicación de qué "
        "hay que corregir.\n\n"
        "Una vez corregido, se vuelve a importar el archivo. Esta hoja puede "
        "dejarse: el sistema sólo lee las hojas que espera."
    )
    celda = guia.cell(row=fila, column=2, value=ayuda)
    celda.alignment = Alignment(wrap_text=True, vertical="top")
    guia.row_dimensions[fila].height = 190

    buffer = io.BytesIO()
    libro.save(buffer)
    base = os.path.splitext(cabecera["nombre_archivo"])[0]
    return buffer.getvalue(), f"{base}_CON_ERRORES_MARCADOS.xlsx"
