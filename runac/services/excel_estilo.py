"""El aspecto común de los Excel que genera el MIR: informes e instructivo.

Colores del Design System de SISOC (los mismos de la pantalla): encabezado en
el verde de la navegación con la franja ámbar, y rojo y ámbar por gravedad, en
sus tonos oscuros para llegar al contraste AA. Letra Segoe UI: es el respaldo
que declara el Design System y está en todas las PC con Windows; Roboto no
suele estar instalada y Excel la reemplaza por otra.

Las solapas NO llevan color. Excel muestra la solapa activa clara y las demás
con su color lleno: una solapa pintada llama más la atención que la que se
está mirando, y parecía que la activa era otra (26-09-2026).

Desde el 28-09-2026 todos llevan la marca —«SISOC · MIR v1.0 · RUNAC»— en la
franja ámbar de arriba y en el pie de impresión, y la plantilla usa la misma
jerarquía de verdes (`motor/plantilla.py`): grupos en el verde más oscuro,
obligatorios en verde, opcionales en verde clarito.
"""

from __future__ import annotations

from datetime import datetime

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.properties import PageSetupProperties

LETRA = "Segoe UI"
NAV, AMBAR, TEAL = "045F5B", "FFC000", "04756F"
TINTA, TINTA2, LINEA, TENUE = "1C1917", "57534E", "E7E5E4", "8A837D"

# Superficie, texto y borde de cada tono.
BLOQUEANTE = {"sup": "FDECEC", "tin": "7F1D1D", "bor": "C62828"}
ADVERTENCIA = {"sup": "FBEEE1", "tin": "7A3B00", "bor": "E86A00"}
INFO = {"sup": "E3F2FB", "tin": "01466E", "bor": "01699F"}
PENDIENTE = {"sup": "EFEFEE", "tin": "44403C", "bor": "78716C"}
BIEN = {"sup": "C6E3E1", "tin": "045F5B", "bor": "04756F"}
POR_GRAVEDAD = {"BLOQUEANTE": BLOQUEANTE, "ADVERTENCIA": ADVERTENCIA}

# Los colores con que la plantilla pinta los títulos: los mismos de
# `motor/plantilla.py`.
VERDE_GRUPO, VERDE_OBLIGATORIO, VERDE_OPCIONAL = "073B38", "9DCECB", "E4F1F0"


def marca() -> str:
    """«SISOC · MIR v1.0 · RUNAC»: ecosistema, módulo e implementación."""
    try:
        from django.conf import settings  # pylint: disable=import-outside-toplevel

        return settings.MIR_MARCA
    except Exception:  # pylint: disable=broad-except  # fuera de Django
        return "SISOC · MIR v1.0 · RUNAC"


def fuente(**k) -> Font:
    return Font(name=LETRA, **k)


def relleno(color: str) -> PatternFill:
    return PatternFill("solid", fgColor=color)


def ahora_legible(v: datetime | None) -> str:
    return v.strftime("%d/%m/%Y %H:%M") if isinstance(v, datetime) else (v or "")


def encabezado(ws, titulos: list[str], anchos: list[int], fila: int = 1) -> None:
    """La fila de títulos de una tabla, con filtros y fija al desplazarse."""
    for i, (titulo, ancho) in enumerate(zip(titulos, anchos), start=1):
        c = ws.cell(row=fila, column=i, value=titulo)
        c.font = fuente(size=10, bold=True, color="FFFFFF")
        c.fill = relleno(NAV)
        c.alignment = Alignment(vertical="center", wrap_text=True)
        c.border = Border(bottom=Side(style="thick", color=AMBAR))
        ws.column_dimensions[c.column_letter].width = ancho
    ws.row_dimensions[fila].height = 30
    # Por coordenada: pedir la celda la crea, y la fila siguiente quedaba vacía.
    ws.freeze_panes = f"A{fila + 1}"


def cuerpo(c, **k) -> None:
    c.font = fuente(size=10, color=TINTA, **k)
    c.alignment = Alignment(vertical="top", wrap_text=True)
    c.border = Border(bottom=Side(style="thin", color=LINEA))


def etiqueta(c, tono: dict) -> None:
    """Una celda pintada como las etiquetas de la pantalla."""
    c.fill = relleno(tono["sup"])
    c.font = fuente(size=10, bold=True, color=tono["tin"])


def alto_para(ws, fila: int, textos, caracteres_por_renglon: int = 48) -> None:
    renglones = max(
        [1]
        + [
            sum(
                max(1, len(parte) // caracteres_por_renglon + 1)
                for parte in str(t or "").split("\n")
            )
            for t in textos
        ]
    )
    ws.row_dimensions[fila].height = 15 * renglones + 2


def para_imprimir(ws, pie: str, repetir: str | None = None) -> None:
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    if repetir:
        ws.print_title_rows = repetir
    ws.oddFooter.left.text = f"{marca()} · {pie}"
    ws.oddFooter.right.text = "Página &P de &N"
    for parte in (ws.oddFooter.left, ws.oddFooter.right):
        parte.size, parte.font = 8, LETRA


def portada(ws, titulo: str, subtitulo: str, para_que: str, ancho: int = 6) -> None:
    """El encabezado de una hoja de resumen: título, de qué se trata y la franja."""
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3
    # La marca, en una franja ámbar arriba de todo: el rastro de dónde salió.
    ws["B1"] = marca()
    ws["B1"].font = fuente(size=9, bold=True, color=TINTA)
    for col in range(2, 2 + ancho):
        ws.cell(row=1, column=col).fill = relleno(AMBAR)
    ws.row_dimensions[1].height = 18
    ws["B2"] = titulo
    ws["B2"].font = fuente(size=18, bold=True, color=TEAL)
    ws["B3"] = subtitulo
    ws["B3"].font = fuente(size=11, color=TINTA2)
    if para_que:
        ws["B4"] = para_que
        ws["B4"].font = fuente(size=10, italic=True, color=TINTA2)
        ws["B4"].alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=4, start_column=2, end_row=4, end_column=1 + ancho)
        ws.row_dimensions[4].height = 30
    for col in range(2, 2 + ancho):
        ws.cell(row=5, column=col).border = Border(
            bottom=Side(style="medium", color=AMBAR)
        )


def tarjetas(ws, fila: int, datos: list[tuple[str, int, dict]]) -> None:
    """Números grandes, uno por columna a partir de la B."""
    for i, (texto, numero, tono) in enumerate(datos, start=2):
        arriba, abajo = ws.cell(row=fila, column=i), ws.cell(row=fila + 1, column=i)
        for c in (arriba, abajo):
            c.fill = relleno(tono["sup"])
            c.alignment = Alignment(horizontal="center", vertical="center")
        arriba.value = numero
        arriba.font = fuente(size=20, bold=True, color=tono["tin"])
        arriba.border = Border(top=Side(style="thick", color=tono["bor"]))
        abajo.value = texto
        abajo.font = fuente(size=9, color=tono["tin"])
    ws.row_dimensions[fila].height = 34
