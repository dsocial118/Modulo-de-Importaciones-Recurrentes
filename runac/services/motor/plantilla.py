"""Genera el Excel modelo que se le entrega a las provincias, leyendo la Capa 1.

    python plantilla.py --archivo MPI --salida /trabajo/capa1/plantillas
    python plantilla.py --todos --periodo 2026_T1

La plantilla se arma con la MISMA configuración contra la que después se valida.
Eso garantiza algo que hoy no está garantizado: que la planilla que completa la
provincia y la que espera el sistema sean exactamente la misma.

Cada archivo sale con:
  - el título y el subtítulo de la planilla;
  - la fila de dimensiones, con las celdas combinadas donde corresponde;
  - los títulos de columna, en el orden definido;
  - las listas desplegables cargadas con los valores vigentes;
  - la explicación de cada campo como comentario en la celda del título: al
    pasar el mouse por encima aparece el globo, sin ir a otra hoja;
  - una hoja de listas con los catálogos, oculta.
"""

from __future__ import annotations

import argparse
import os
from datetime import date

import mysql.connector
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

CONEXION = dict(
    host=os.environ.get("RUNAC_DB_HOST", "mysql"),
    port=3306,
    user="root",
    password="runac_local",
    database="runac",
)

# Excel no admite más de 255 caracteres en una lista escrita dentro de la
# validación. Por encima de eso, la lista tiene que vivir en otra hoja.
LIMITE_LISTA_INLINE = 250

AZUL = "1F4E79"
AZUL_CLARO = "DDEBF7"
GRIS = "F2F2F2"
BORDE = Border(*[Side(style="thin", color="BFBFBF")] * 4)


def leer_definicion(cur, codigo: str) -> dict:
    cur.execute(
        """
        SELECT a.id, a.codigo, a.descripcion,
               av.id AS version_id, av.numero AS version,
               av.nombre_esperado, av.titulo, av.subtitulo
        FROM runac_c1_archivo a
        JOIN runac_c1_archivo_version av ON av.archivo_id = a.id AND av.estado = 'VIGENTE'
        WHERE a.codigo = %s
    """,
        (codigo,),
    )
    archivo = cur.fetchone()
    if not archivo:
        raise SystemExit(
            f"No existe el archivo {codigo} con una versión vigente en la Capa 1."
        )

    cur.execute(
        """
        SELECT id, nombre_esperado, descripcion, orden_procesamiento, fila_encabezados
        FROM runac_c1_hoja WHERE archivo_version_id = %s ORDER BY orden_procesamiento
    """,
        (archivo["version_id"],),
    )
    hojas = cur.fetchall()

    for h in hojas:
        cur.execute(
            """
            SELECT id, nombre_esperado, orden FROM runac_c1_dimension
            WHERE hoja_id = %s ORDER BY orden
        """,
            (h["id"],),
        )
        h["dimensiones"] = cur.fetchall()

        cur.execute(
            """
            SELECT c.id, c.nombre, c.titulo_esperado, c.orden, c.tipo_dato,
                   c.longitud_maxima, c.obligatorio, c.ayuda,
                   d.nombre_esperado AS dimension, cat.codigo AS catalogo, cat.nombre AS catalogo_nombre
            FROM runac_c1_campo c
            LEFT JOIN runac_c1_dimension d ON d.id = c.dimension_id
            LEFT JOIN runac_c1_catalogo cat ON cat.id = c.catalogo_id
            WHERE c.hoja_id = %s ORDER BY c.orden
        """,
            (h["id"],),
        )
        h["campos"] = cur.fetchall()

    # Los catálogos que efectivamente usa este archivo.
    cur.execute(
        """
        SELECT DISTINCT cat.codigo, cat.nombre
        FROM runac_c1_campo c
        JOIN runac_c1_hoja h ON h.id = c.hoja_id
        JOIN runac_c1_catalogo cat ON cat.id = c.catalogo_id
        WHERE h.archivo_version_id = %s ORDER BY cat.codigo
    """,
        (archivo["version_id"],),
    )
    catalogos = cur.fetchall()
    for cat in catalogos:
        cur.execute(
            """
            SELECT o.valor_esperado FROM runac_c1_catalogo_opcion o
            JOIN runac_c1_catalogo c ON c.id = o.catalogo_id
            WHERE c.codigo = %s AND o.activo = 1 ORDER BY o.orden
        """,
            (cat["codigo"],),
        )
        cat["valores"] = [r["valor_esperado"] for r in cur.fetchall()]

    archivo["hojas"] = hojas
    archivo["catalogos"] = catalogos
    return archivo


TIPOS_EN_CASTELLANO = {
    "FECHA": "una fecha, con formato dd/mm/aaaa",
    "ENTERO": "un número entero",
    "DECIMAL": "un número, puede tener decimales",
    "TEXTO": "texto",
}


def _comentario_del_campo(campo: dict) -> Comment:
    """El globo que aparece al pasar el mouse sobre el título de la columna.

    Reúne todo lo que necesita saber quien completa la planilla: qué se espera,
    de qué tipo, si es obligatorio y si hay que elegir de una lista.
    """
    partes = [campo["titulo_esperado"], ""]

    if campo["ayuda"]:
        partes.append(campo["ayuda"])
        partes.append("")

    detalle = ["Qué se espera: " + TIPOS_EN_CASTELLANO.get(campo["tipo_dato"], "texto")]
    if campo["tipo_dato"] == "TEXTO" and campo["longitud_maxima"]:
        detalle.append(f'Máximo {campo["longitud_maxima"]} caracteres.')
    if campo["catalogo_nombre"]:
        detalle.append(f'Elegir de la lista "{campo["catalogo_nombre"]}".')
    detalle.append(
        "Campo obligatorio." if campo["obligatorio"] else "No es obligatorio."
    )
    partes.extend(detalle)

    texto = "\n".join(partes)
    comentario = Comment(texto, "RUNAC")
    # El globo se dimensiona según el largo del texto, para que se lea entero.
    lineas = sum(max(1, len(linea) // 48 + 1) for linea in texto.split("\n"))
    comentario.width = 340
    comentario.height = min(400, max(90, lineas * 16 + 20))
    return comentario


def escribir_hoja_listas(wb, catalogos) -> dict[str, str]:
    """Escribe una hoja con los catálogos y devuelve el rango de cada uno."""
    ws = wb.create_sheet("LISTAS")
    rangos = {}
    for i, cat in enumerate(catalogos, start=1):
        letra = get_column_letter(i)
        ws.cell(row=1, column=i, value=cat["nombre"][:255]).font = Font(
            bold=True, color="FFFFFF"
        )
        ws.cell(row=1, column=i).fill = PatternFill("solid", fgColor=AZUL)
        for j, v in enumerate(cat["valores"], start=2):
            ws.cell(row=j, column=i, value=v)
        ws.column_dimensions[letra].width = 28
        if cat["valores"]:
            rangos[cat["codigo"]] = (
                f"LISTAS!${letra}$2:${letra}${len(cat['valores']) + 1}"
            )
    ws.sheet_state = "visible"
    ws["A1"].comment = None
    return rangos


def escribir_hoja_datos(wb, archivo, hoja, rangos, filas_vacias: int):
    ws = wb.create_sheet(hoja["nombre_esperado"][:31])
    campos = hoja["campos"]
    n = len(campos)
    if n == 0:
        return

    # La fila de encabezados la manda la Capa 1: es contra ese número que después
    # se valida el archivo que sube la provincia. Todo lo demás se acomoda arriba.
    fila_enc = hoja["fila_encabezados"]
    fila_dim = (fila_enc - 1) if hoja["dimensiones"] and fila_enc > 1 else None
    fila_titulo = (
        (fila_dim - 1) if fila_dim else (fila_enc - 1 if fila_enc > 1 else None)
    )

    if archivo["titulo"] and fila_titulo and fila_titulo >= 1:
        ws.cell(row=fila_titulo, column=1, value=archivo["titulo"])
        ws.merge_cells(
            start_row=fila_titulo, start_column=1, end_row=fila_titulo, end_column=n
        )
        c = ws.cell(row=fila_titulo, column=1)
        c.font = Font(bold=True, size=14, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=AZUL)
        c.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[fila_titulo].height = 24
    if archivo["subtitulo"] and fila_titulo and fila_titulo > 1:
        ws.cell(row=fila_titulo - 1, column=1, value=archivo["subtitulo"])
        ws.merge_cells(
            start_row=fila_titulo - 1,
            start_column=1,
            end_row=fila_titulo - 1,
            end_column=n,
        )
        ws.cell(row=fila_titulo - 1, column=1).font = Font(italic=True, size=11)
        ws.cell(row=fila_titulo - 1, column=1).alignment = Alignment(
            horizontal="center"
        )

    # Fila de dimensiones: se combinan las columnas contiguas de cada una.
    if fila_dim:
        i = 0
        while i < n:
            dim = campos[i]["dimension"]
            j = i
            while j + 1 < n and campos[j + 1]["dimension"] == dim:
                j += 1
            if dim:
                ws.cell(row=fila_dim, column=i + 1, value=dim)
                if j > i:
                    ws.merge_cells(
                        start_row=fila_dim,
                        start_column=i + 1,
                        end_row=fila_dim,
                        end_column=j + 1,
                    )
                c = ws.cell(row=fila_dim, column=i + 1)
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = PatternFill("solid", fgColor="2E75B6")
                c.alignment = Alignment(horizontal="center", vertical="center")
            i = j + 1

    for k, campo in enumerate(campos, start=1):
        titulo = campo["titulo_esperado"]
        if campo["obligatorio"]:
            titulo += " *"
        c = ws.cell(row=fila_enc, column=k, value=titulo)
        c.font = Font(bold=True)
        c.fill = PatternFill(
            "solid", fgColor=AZUL_CLARO if campo["obligatorio"] else GRIS
        )
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        c.border = BORDE

        # La explicación va en la propia celda del título, como comentario: al
        # pasar el mouse por encima aparece el globo, sin tener que ir a otra hoja.
        c.comment = _comentario_del_campo(campo)

        ancho = min(40, max(14, len(campo["titulo_esperado"]) // 2 + 8))
        ws.column_dimensions[get_column_letter(k)].width = ancho
    ws.row_dimensions[fila_enc].height = 48
    ws.freeze_panes = ws.cell(row=fila_enc + 1, column=1)

    primera_dato = fila_enc + 1
    ultima_dato = fila_enc + filas_vacias

    # Formato y validaciones por columna.
    for k, campo in enumerate(campos, start=1):
        letra = get_column_letter(k)
        rango = f"{letra}{primera_dato}:{letra}{ultima_dato}"

        if campo["tipo_dato"] == "FECHA":
            for r in range(primera_dato, ultima_dato + 1):
                ws.cell(row=r, column=k).number_format = "dd/mm/yyyy"
        elif campo["tipo_dato"] in ("ENTERO", "DECIMAL"):
            fmt = "0" if campo["tipo_dato"] == "ENTERO" else "0.00"
            for r in range(primera_dato, ultima_dato + 1):
                ws.cell(row=r, column=k).number_format = fmt

        if campo["catalogo"] and campo["catalogo"] in rangos:
            # allowBlank explícito: dice si el campo admite quedar vacío. Es la
            # marca que después el importador lee como obligatoriedad.
            dv = DataValidation(
                type="list",
                formula1=rangos[campo["catalogo"]],
                allow_blank=not campo["obligatorio"],
                showDropDown=False,
            )
            dv.error = f'El valor no está entre los admitidos para "{campo["titulo_esperado"]}".'
            dv.errorTitle = "Valor no admitido"
            if campo["ayuda"]:
                dv.prompt = campo["ayuda"][:255]
                dv.promptTitle = campo["titulo_esperado"][:32]
            ws.add_data_validation(dv)
            dv.add(rango)

    return {"fila_encabezados": fila_enc, "filas": (primera_dato, ultima_dato)}


def escribir_hoja_instrucciones(wb, archivo):
    ws = wb.create_sheet("INSTRUCCIONES")
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 95
    ws.cell(row=1, column=1, value="Campo").font = Font(bold=True, color="FFFFFF")
    ws.cell(row=1, column=2, value="Indicación para el llenado").font = Font(
        bold=True, color="FFFFFF"
    )
    for c in ("A1", "B1"):
        ws[c].fill = PatternFill("solid", fgColor=AZUL)

    f = 2
    for hoja in archivo["hojas"]:
        ws.cell(row=f, column=1, value=f'Hoja: {hoja["nombre_esperado"]}').font = Font(
            bold=True, size=12
        )
        f += 1
        dimension_actual = object()
        for campo in hoja["campos"]:
            if campo["dimension"] != dimension_actual:
                dimension_actual = campo["dimension"]
                if dimension_actual:
                    c = ws.cell(row=f, column=1, value=dimension_actual)
                    c.font = Font(bold=True)
                    c.fill = PatternFill("solid", fgColor=AZUL_CLARO)
                    f += 1
            ws.cell(row=f, column=1, value=campo["titulo_esperado"])
            partes = []
            if campo["ayuda"]:
                partes.append(campo["ayuda"])
            if campo["obligatorio"]:
                partes.append("Campo obligatorio.")
            if campo["catalogo_nombre"]:
                partes.append(
                    f'Debe elegirse un valor de la lista "{campo["catalogo_nombre"]}".'
                )
            if campo["tipo_dato"] == "FECHA":
                partes.append("Formato de fecha: dd/mm/aaaa.")
            ws.cell(row=f, column=2, value="\n".join(partes) or None).alignment = (
                Alignment(wrap_text=True, vertical="top")
            )
            f += 1
        f += 1


def generar(
    cur, codigo: str, salida: str, filas_vacias: int, periodo: str | None
) -> str:
    archivo = leer_definicion(cur, codigo)
    wb = Workbook()
    wb.remove(wb.active)

    rangos = escribir_hoja_listas(wb, archivo["catalogos"])
    for hoja in archivo["hojas"]:
        escribir_hoja_datos(wb, archivo, hoja, rangos, filas_vacias)

    # Hoja con la ayuda de cada campo: es lo que le explica al operador que se
    # espera en cada columna, sin tener que consultar el documento funcional.
    escribir_hoja_instrucciones(wb, archivo)

    # La hoja de listas va al final y oculta: ayuda al usuario sin estorbarlo.
    wb.move_sheet("LISTAS", offset=len(wb.sheetnames))
    wb["LISTAS"].sheet_state = "hidden"

    wb.properties.title = archivo["titulo"] or codigo
    wb.properties.subject = f"RUNAC — plantilla {codigo}"
    wb.properties.description = (
        f"Generada desde la Capa 1 el {date.today().isoformat()}"
        + (f" para el período {periodo}." if periodo else ".")
        + " No modificar los títulos de las columnas ni su orden."
    )

    sufijo = f"_{periodo}" if periodo else ""
    destino = os.path.join(salida, f"{codigo}{sufijo}_MODELO.xlsx")
    os.makedirs(salida, exist_ok=True)
    wb.save(destino)
    return destino


def main():
    p = argparse.ArgumentParser(
        description="Genera las plantillas Excel desde la Capa 1."
    )
    p.add_argument("--archivo", default=None)
    p.add_argument("--todos", action="store_true")
    p.add_argument("--salida", default="/trabajo/capa1/plantillas")
    p.add_argument(
        "--filas", type=int, default=200, help="filas vacías preparadas para cargar"
    )
    p.add_argument("--periodo", default=None)
    args = p.parse_args()

    cn = mysql.connector.connect(**CONEXION)
    cur = cn.cursor(dictionary=True)

    if args.todos:
        cur.execute(
            """SELECT a.codigo FROM runac_c1_archivo a
                       JOIN runac_c1_archivo_version av ON av.archivo_id = a.id AND av.estado='VIGENTE'
                       ORDER BY av.orden_importacion"""
        )
        codigos = [r["codigo"] for r in cur.fetchall()]
    elif args.archivo:
        codigos = [args.archivo]
    else:
        raise SystemExit("Indicar --archivo CODIGO o --todos.")

    for codigo in codigos:
        destino = generar(cur, codigo, args.salida, args.filas, args.periodo)
        tam = os.path.getsize(destino) / 1024
        print(f"  {codigo:12} {os.path.basename(destino):34} {tam:7.0f} KB")

    cur.close()
    cn.close()


if __name__ == "__main__":
    main()
